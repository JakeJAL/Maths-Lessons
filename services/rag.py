"""RAG system — index example slides and search for relevant curriculum content."""

import os
import hashlib
import chromadb
from pptx import Presentation

EXAMPLE_SLIDES_DIR = "example slides"
CHROMA_PERSIST_DIR = ".chroma_db"


def _get_collection():
    """Get or create the ChromaDB collection."""
    client = chromadb.PersistentClient(path=CHROMA_PERSIST_DIR)
    return client.get_or_create_collection(
        name="curriculum_slides",
        metadata={"hnsw:space": "cosine"},
    )


def _extract_slidedeck_text(pptx_path: str) -> list[dict]:
    """Extract text from a .pptx file, returning a list of slide dicts."""
    try:
        prs = Presentation(pptx_path)
    except Exception as e:
        print(f"[RAG] Failed to open {pptx_path}: {e}")
        return []

    slides = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t and len(t) > 2:  # skip tiny labels
                    texts.append(t)
        if texts:
            slides.append({
                "index": i,
                "text": "\n".join(texts),
            })
    return slides


def _parse_topic_and_lesson(pptx_path: str) -> tuple[str, str]:
    """Extract topic and lesson name from the file path."""
    parts = pptx_path.replace("\\", "/").split("/")
    # Structure: example slides/<topic>/<lesson>/slidedeck.pptx
    topic = ""
    lesson = ""
    for i, p in enumerate(parts):
        if p == "example slides" and i + 1 < len(parts):
            topic = parts[i + 1].replace("-", " ").rstrip("0123456789 ")
            if i + 2 < len(parts):
                lesson = parts[i + 2].replace("-", " ").lstrip("0123456789 ")
            break
    return topic.strip(), lesson.strip()


def index_example_slides():
    """Scan example slides folder and index all slidedecks into ChromaDB."""
    collection = _get_collection()

    # Check if already indexed
    existing = collection.count()
    if existing > 0:
        print(f"[RAG] Collection already has {existing} documents, skipping indexing.")
        return existing

    print("[RAG] Indexing example slides...")
    doc_count = 0
    ids_batch = []
    docs_batch = []
    metas_batch = []

    for root, dirs, files in os.walk(EXAMPLE_SLIDES_DIR):
        for fname in files:
            # Only index main slidedecks, not worksheets
            if fname != "slidedeck.pptx":
                continue

            fpath = os.path.join(root, fname)
            topic, lesson = _parse_topic_and_lesson(fpath)
            if not topic:
                continue

            slides = _extract_slidedeck_text(fpath)
            if not slides:
                continue

            # One document per lesson (full text, trimmed)
            full_text = f"Topic: {topic}\nLesson: {lesson}\n\n"
            for s in slides:
                full_text += s["text"] + "\n"

            doc_id = hashlib.md5(fpath.encode()).hexdigest()[:12]
            ids_batch.append(f"{doc_id}_full")
            docs_batch.append(full_text[:5000])
            metas_batch.append({
                "topic": topic,
                "lesson": lesson,
                "file": fpath,
                "type": "full_lesson",
                "slide_count": len(slides),
            })
            doc_count += 1

            # Batch insert every 50 docs
            if len(ids_batch) >= 50:
                collection.add(ids=ids_batch, documents=docs_batch, metadatas=metas_batch)
                print(f"[RAG] Indexed {doc_count} lessons...")
                ids_batch, docs_batch, metas_batch = [], [], []

    # Final batch
    if ids_batch:
        collection.add(ids=ids_batch, documents=docs_batch, metadatas=metas_batch)

    print(f"[RAG] Indexed {doc_count} lessons.")
    return doc_count


def search_curriculum(query: str, n_results: int = 5) -> list[dict]:
    """Search the curriculum content for relevant lessons.

    Returns a list of {topic, lesson, content, relevance} dicts.
    """
    collection = _get_collection()

    if collection.count() == 0:
        print("[RAG] Collection is empty, nothing to search.")
        return []

    results = collection.query(
        query_texts=[query],
        n_results=n_results,
    )

    matches = []
    if results and results["documents"]:
        for i, doc in enumerate(results["documents"][0]):
            meta = results["metadatas"][0][i] if results["metadatas"] else {}
            distance = results["distances"][0][i] if results["distances"] else None
            matches.append({
                "topic": meta.get("topic", ""),
                "lesson": meta.get("lesson", ""),
                "content": doc[:1500],
                "slide_count": meta.get("slide_count", 0),
                "relevance": round(1 - (distance or 0), 3),
            })

    return matches

    return matches

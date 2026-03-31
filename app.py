import os
import re
from flask import Flask, render_template, request, jsonify, session, redirect, url_for
from agents.orchestrator import run_agent
from tools import execute_tool
from state import state, reset_state
from services.auth import is_logged_in, get_login_url, handle_callback, logout

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", os.urandom(32).hex())

def _classify_shape_role(text: str, shape) -> str:
    """Classify a text shape's role based on its content."""
    from services.slides_api import _normalize_math_text
    t = _normalize_math_text(text).strip().lower()
    import re as _re

    # Detect equations: contains math operators/variables in a pattern like "5x + 11 = 29"
    if _re.search(r'\d+[a-z]|\d+\s*[+\-*/÷×]\s*\d+|=\s*\d+|\\frac', t):
        return "equation"
    # Labels like "Example 1", "Question 3", "Task 2"
    if _re.match(r'^(example|question|task|exercise|problem|worked example|practice)\s*\d*\s*$', t):
        return "label"
    # Instructions like "Solve:", "Calculate:", "Work out:", "Simplify:"
    if _re.match(r'^(solve|calculate|work out|simplify|evaluate|find|expand|factorise|sketch|show that|prove)\s*:?\s*$', t):
        return "instruction"
    # Step-by-step solutions (numbered steps or "Step 1:" patterns)
    if _re.search(r'^\d+[\.\)]\s', t) or _re.search(r'^step\s*\d', t):
        return "steps"
    # Answer patterns
    if t.startswith("answer") or t.startswith("ans ") or t.startswith("ans:") or _re.match(r'^[a-z]\s*=\s*\d', t):
        return "answer"
    # Title-like: short text, often the slide heading
    if len(t) < 40 and '\n' not in t:
        return "heading"
    # Default: body content
    return "body"


SLIDES_URL_PATTERN = re.compile(
    r'\[([^\]]*)\]\(https://docs\.google\.com/presentation/d/[^)]+\)'
)
SLIDES_RAW_URL_PATTERN = re.compile(
    r'https://docs\.google\.com/presentation/d/\S+'
)


def _rgb_from_color(color):
    """Extract RGB dict from a pptx RGBColor or theme color."""
    try:
        if color and color.rgb:
            r, g, b = color.rgb
            return {"red": round(r / 255, 3), "green": round(g / 255, 3), "blue": round(b / 255, 3)}
    except (AttributeError, TypeError):
        pass
    return None


def _extract_presentation_style(prs):
    """Extract visual style info from a pptx Presentation object."""
    from pptx.util import Pt, Emu
    from pptx.dml.color import RGBColor

    style = {
        "background_color": None,
        "heading_font_size": None,
        "heading_color": None,
        "heading_bold": None,
        "body_font_size": None,
        "body_color": None,
        "slide_width_pt": None,
        "slide_height_pt": None,
    }

    # Slide dimensions
    if prs.slide_width:
        style["slide_width_pt"] = round(prs.slide_width / 12700, 1)
    if prs.slide_height:
        style["slide_height_pt"] = round(prs.slide_height / 12700, 1)

    # Sample styles from the first few slides
    heading_sizes = []
    heading_colors = []
    heading_bolds = []
    body_sizes = []
    body_colors = []
    bg_colors = []

    for slide in list(prs.slides)[:5]:
        # Background
        bg = slide.background
        if bg and bg.fill and bg.fill.type is not None:
            try:
                fc = bg.fill.fore_color
                rgb = _rgb_from_color(fc)
                if rgb:
                    bg_colors.append(rgb)
            except Exception:
                pass

        # Text styles from shapes
        text_shapes = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                text_shapes.append(shape)

        for idx, shape in enumerate(text_shapes[:2]):
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    size = run.font.size
                    bold = run.font.bold
                    color = _rgb_from_color(run.font.color)

                    if idx == 0:  # heading
                        if size:
                            heading_sizes.append(round(size / 12700, 1))
                        if color:
                            heading_colors.append(color)
                        if bold is not None:
                            heading_bolds.append(bold)
                    else:  # body
                        if size:
                            body_sizes.append(round(size / 12700, 1))
                        if color:
                            body_colors.append(color)

    # Pick most common values
    if heading_sizes:
        style["heading_font_size"] = max(set(heading_sizes), key=heading_sizes.count)
    if heading_colors:
        style["heading_color"] = heading_colors[0]
    if heading_bolds:
        style["heading_bold"] = max(set(heading_bolds), key=heading_bolds.count)
    if body_sizes:
        style["body_font_size"] = max(set(body_sizes), key=body_sizes.count)
    if body_colors:
        style["body_color"] = body_colors[0]
    if bg_colors:
        style["background_color"] = bg_colors[0]

    # Strip None values
    return {k: v for k, v in style.items() if v is not None}


def _extract_shape_text(shape, etree):
    """Extract all text from a shape, including math equations and tables."""
    parts = []

    # Regular text frames (also dig into XML for math)
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            para_parts = []
            # Get regular text runs
            for run in para.runs:
                if run.text.strip():
                    para_parts.append(run.text.strip())
            # Look for math elements in the paragraph XML
            math_els = para._p.findall(
                ".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
            )
            for math_el in math_els:
                math_text = "".join(math_el.itertext()).strip()
                if math_text:
                    para_parts.append(math_text)
            if para_parts:
                parts.append(" ".join(para_parts))

    # Also check the raw shape XML for math outside text frames
    if hasattr(shape, '_element'):
        math_els = shape._element.findall(
            ".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
        )
        for math_el in math_els:
            math_text = "".join(math_el.itertext()).strip()
            if math_text and math_text not in " ".join(parts):
                parts.append(math_text)

    # Tables
    if shape.has_table:
        for row in shape.table.rows:
            row_texts = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_texts.append(cell_text)
            if row_texts:
                parts.append(" | ".join(row_texts))

    # Grouped shapes — recurse
    if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
        for child in shape.shapes:
            child_text = _extract_shape_text(child, etree)
            if child_text:
                parts.append(child_text)

    return "\n".join(parts)


@app.route("/")
def index():
    if not is_logged_in():
        return redirect(url_for("login"))
    return render_template("index.html",
        user_name=session.get("user_name", ""),
        user_email=session.get("user_email", ""),
    )


@app.route("/login")
def login():
    if is_logged_in():
        return redirect(url_for("index"))
    return render_template("login.html")


@app.route("/auth/start")
def auth_start():
    return redirect(get_login_url())


@app.route("/auth/callback")
def auth_callback():
    try:
        handle_callback()
        return redirect(url_for("index"))
    except Exception as e:
        return f"Authentication failed: {e}", 400


@app.route("/auth/logout")
def auth_logout():
    logout()
    return redirect(url_for("login"))


@app.route("/api/chat", methods=["POST"])
def chat():
    if not is_logged_in():
        return jsonify({"error": "Not authenticated"}), 401
    data = request.get_json()
    user_message = data.get("message", "").strip()
    if not user_message:
        return jsonify({"error": "Empty message"}), 400

    state["messages"].append({"role": "user", "content": user_message})

    history = [
        {"role": m["role"], "content": m["content"]}
        for m in state["messages"][:-1]
    ]
    pres_context = {
        "id": state["presentation_id"],
        "title": state["presentation_title"],
        "slides": state["current_slides"],
        "uploaded_references": state["uploaded_references"],
    }

    try:
        reply, tool_log = run_agent(
            user_message, history, pres_context, execute_tool
        )
        reply = reply or ""

        # Remove any hallucinated slides URLs
        reply = SLIDES_URL_PATTERN.sub("", reply).strip()
        reply = SLIDES_RAW_URL_PATTERN.sub("", reply).strip()

        # Add the real link when the agent used tools
        if tool_log and state["presentation_id"]:
            pid = state["presentation_id"]
            link = "https://docs.google.com/presentation/d/" + pid
            reply = reply + "\n\n[Open in Google Slides](" + link + ")"

        state["messages"].append({"role": "assistant", "content": reply})
        return jsonify({"reply": reply})

    except Exception as e:
        import traceback
        traceback.print_exc()
        msg = "Something went wrong: " + str(e)
        state["messages"].append({"role": "assistant", "content": msg})
        return jsonify({"reply": msg}), 500


@app.route("/api/reset", methods=["POST"])
def reset():
    reset_state()
    return jsonify({"status": "ok"})


@app.route("/api/upload", methods=["POST"])
def upload():
    if not is_logged_in():
        return jsonify({"error": "Not authenticated"}), 401
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400

    file = request.files["file"]
    if not file.filename.endswith((".pptx", ".pptm")):
        return jsonify({"error": "Only .pptx and .pptm files are supported"}), 400

    try:
        from pptx import Presentation
        from lxml import etree
        import tempfile

        # Save to a temp file for Drive upload
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=file.filename)
        file.save(tmp.name)
        tmp.close()

        # Parse content locally
        prs = Presentation(tmp.name)
        slides = []
        style = _extract_presentation_style(prs)

        for i, slide in enumerate(prs.slides):
            texts = []
            text_elements = []  # structured per-shape info
            for shape in slide.shapes:
                extracted = _extract_shape_text(shape, etree)
                if extracted:
                    texts.append(extracted)
                # Record text box details for shape-level replacement
                if shape.has_text_frame and shape.text_frame.text.strip():
                    shape_text = shape.text_frame.text.strip()
                    role = _classify_shape_role(shape_text, shape)
                    text_elements.append({
                        "index": len(text_elements),
                        "role": role,
                        "text": shape_text,
                    })
            slides.append({
                "index": i,
                "content": "\n".join(texts),
                "text_elements": text_elements,
            })

        # Upload to Google Drive as a Slides template
        from services.slides_api import upload_pptx_as_template
        template_id = upload_pptx_as_template(tmp.name, file.filename)

        # Clean up temp file
        os.unlink(tmp.name)

        ref = {
            "filename": file.filename,
            "slide_count": len(slides),
            "slides": slides,
            "style": style,
            "template_id": template_id,
        }
        state["uploaded_references"].append(ref)

        return jsonify({
            "status": "ok",
            "filename": file.filename,
            "slide_count": len(slides),
            "template_ready": True,
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


# Index RAG on first import (already persisted after Docker build, so this is fast)
import os as _os
if not _os.environ.get("SKIP_RAG_INDEX"):
    from services.rag import index_example_slides
    index_example_slides()


if __name__ == "__main__":
    print("Starting server...")
    # Allow OAuth over HTTP for local dev
    os.environ["OAUTHLIB_INSECURE_TRANSPORT"] = "1"
    app.run(debug=True, port=5000)
    app.run(debug=True, port=5000)
